import pytest
from openpyxl import load_workbook

from app.models import HotelOffer, SearchRequest
from app.services.excel_exporter import export_offers_to_excel


def make_request() -> SearchRequest:
    return SearchRequest.from_mapping(
        {
            "destination": "Kyoto, Japan",
            "check_in": "2026-12-10",
            "check_out": "2026-12-15",
            "adults": 2,
            "rooms": 1,
            "currency": "GHS",
        }
    )


def test_exports_only_requested_columns_and_every_offer(tmp_path) -> None:
    output = tmp_path / "results" / "hotels.xlsx"
    offers = [
        HotelOffer(
            f"River House {number}",
            "Garden room",
            800 + number,
            "GHS",
            8.7,
            address=f"{number} River Road, Kyoto",
            latitude=35.0116 + number / 1000,
            longitude=135.7681 + number / 1000,
        )
        for number in range(12)
    ]

    result = export_offers_to_excel(offers, make_request(), output)

    workbook = load_workbook(result, data_only=True)
    offer_sheet = workbook["Hotel offers"]
    assert offer_sheet.freeze_panes == "A2"
    assert offer_sheet.auto_filter.ref == "A1:H13"
    assert workbook.sheetnames == ["Hotel offers"]
    assert [cell.value for cell in offer_sheet[1]] == [
        "Hotel",
        "Amount per night",
        "Total amount",
        "Currency",
        "Guest rating",
        "Address",
        "Latitude",
        "Longitude",
    ]
    assert [cell.value for cell in offer_sheet[2]] == [
        "River House 0",
        800,
        4000,
        "GHS",
        8.7,
        "0 River Road, Kyoto",
        35.0116,
        135.7681,
    ]
    assert offer_sheet["B2"].number_format == "#,##0.00"
    assert offer_sheet["C2"].number_format == "#,##0.00"
    assert offer_sheet["E2"].number_format == "0.0"
    assert offer_sheet["G2"].number_format == "0.000000"
    assert offer_sheet.max_row == len(offers) + 1
    assert offer_sheet["A13"].value == "River House 11"


def test_exports_empty_results_with_headers_and_creates_parent_directory(tmp_path) -> None:
    output = tmp_path / "nested" / "empty.xlsx"

    export_offers_to_excel([], make_request(), output)

    workbook = load_workbook(output, data_only=True)
    assert workbook["Hotel offers"].max_row == 1
    assert [cell.value for cell in workbook["Hotel offers"][1]] == [
        "Hotel",
        "Amount per night",
        "Total amount",
        "Currency",
        "Guest rating",
        "Address",
        "Latitude",
        "Longitude",
    ]


def test_rejects_non_xlsx_path(tmp_path) -> None:
    with pytest.raises(ValueError, match="must end with \\.xlsx"):
        export_offers_to_excel([], make_request(), tmp_path / "results.csv")
